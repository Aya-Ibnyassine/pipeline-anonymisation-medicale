import streamlit as st
import pandas as pd
import io
import docx
import re
import pdfplumber
import fitz  # PyMuPDF
from agent import AnonymizationAgent 
from pseudonymizer import PatientPseudonymizer 

# 1. Configuration de la page
st.set_page_config(
    page_title="HOSIX | Pipeline Anonymisation", 
    layout="wide",
    initial_sidebar_state="expanded"
)

# 2. Design system moderne (CSS)
st.markdown(
    """
    <style>
    :root {
        --bg: #f5f7fb;
        --surface: rgba(255,255,255,0.86);
        --surface-solid: #ffffff;
        --primary: #0f766e;
        --primary-dark: #115e59;
        --accent: #2563eb;
        --text: #0f172a;
        --muted: #64748b;
        --border: #dbe4f0;
    }

    .stApp {
        background:
            radial-gradient(circle at top left, rgba(37, 99, 235, 0.10), transparent 34%),
            radial-gradient(circle at top right, rgba(15, 118, 110, 0.10), transparent 28%),
            linear-gradient(180deg, #f8fbff 0%, #f5f7fb 100%);
        color: var(--text);
    }

    .block-container {
        padding-top: 1.5rem;
        padding-bottom: 2rem;
    }

    /* === NOUVEL EN-TÊTE HERO (BANNIÈRE CENTRALE) === */
    .hero {
        background: linear-gradient(135deg, rgba(15, 23, 42, 0.96), rgba(15, 118, 110, 0.96));
        color: white;
        border: 1px solid rgba(255,255,255,0.08);
        border-radius: 24px;
        padding: 3.5rem 2rem; /* Plus d'espace vertical pour faire "bannière" */
        text-align: center; /* Centrage du texte */
        box-shadow: 0 18px 50px rgba(15, 23, 42, 0.15);
        margin-bottom: 2rem;
    }

    .hero-title {
        font-size: 2.8rem; /* Titre plus grand */
        font-weight: 800;
        line-height: 1.2;
        margin: 0 0 1rem 0;
        letter-spacing: -0.02em;
    }

    .hero-subtitle {
        color: rgba(255,255,255,0.9);
        font-size: 1.1rem;
        line-height: 1.65;
        max-width: 800px; /* Limite la largeur pour une belle lecture */
        margin: 0 auto; /* Centre le bloc de texte au milieu */
    }
    /* ============================================== */

    /* Styles pour les Popovers (Cartes cliquables) */
    div[data-testid="stPopover"] > button {
        background: rgba(255,255,255,0.85) !important;
        border: 1px solid rgba(219, 228, 240, 0.95) !important;
        border-radius: 16px !important;
        padding: 1rem !important;
        box-shadow: 0 6px 15px rgba(15, 23, 42, 0.04) !important;
        text-align: center !important;
        color: var(--text) !important;
        font-weight: 700 !important;
        font-size: 0.95rem !important;
        transition: all 0.2s ease !important;
        min-height: 70px;
    }
    div[data-testid="stPopover"] > button:hover {
        background: #ffffff !important;
        transform: translateY(-2px) !important;
        box-shadow: 0 10px 25px rgba(15, 23, 42, 0.08) !important;
        border-color: var(--accent) !important;
    }

    /* Cartes (Text, Input) */
    .glass-card, .input-card, .upload-card {
        background: var(--surface);
        border: 1px solid rgba(219, 228, 240, 0.95);
        border-radius: 22px;
        box-shadow: 0 10px 30px rgba(15, 23, 42, 0.06);
        backdrop-filter: blur(10px);
    }

    .text-card {
        background: var(--surface-solid);
        padding: 1.25rem 1.3rem;
        border-radius: 20px;
        border: 1px solid var(--border);
        box-shadow: 0 10px 30px rgba(15, 23, 42, 0.05);
        min-height: 420px;
        font-size: 0.95rem;
        line-height: 1.85;
        color: #334155;
        white-space: pre-wrap;
        font-family: "Inter", "Segoe UI", Arial, sans-serif;
    }

    /* Badges & Boutons d'action */
    .stButton > button {
        background: linear-gradient(135deg, var(--primary), var(--accent));
        color: white !important;
        border: none;
        border-radius: 14px;
        padding: 0.7rem 1.2rem;
        font-weight: 700;
        box-shadow: 0 10px 22px rgba(37, 99, 235, 0.16);
        transition: all 0.2s ease-in-out;
    }
    .stButton > button:hover {
        transform: translateY(-1px);
        box-shadow: 0 14px 24px rgba(37, 99, 235, 0.22);
    }

    .badge-anonyme {
        background: linear-gradient(135deg, rgba(37, 99, 235, 0.12), rgba(15, 118, 110, 0.12));
        color: var(--primary-dark);
        padding: 0.18rem 0.55rem;
        border-radius: 999px;
        font-weight: 800;
        font-size: 0.84em;
        border: 1px solid rgba(37, 99, 235, 0.18);
        white-space: nowrap;
    }

    /* Customisation de la Sidebar et Workflow */
    section[data-testid="stSidebar"] {
        background: linear-gradient(180deg, rgba(15, 23, 42, 0.98), rgba(15, 118, 110, 0.96));
    }
    section[data-testid="stSidebar"] * {
        color: white;
    }
    .sidebar-badge {
        display: inline-block;
        border-radius: 999px;
        padding: 0.28rem 0.75rem;
        font-size: 0.78rem;
        font-weight: 800;
        background: rgba(255,255,255,0.12);
        border: 1px solid rgba(255,255,255,0.14);
        margin-top: 0.35rem;
        margin-bottom: 0.9rem;
    }
    .workflow-box {
        background: rgba(255,255,255,0.08);
        border: 1px solid rgba(255,255,255,0.15);
        border-radius: 14px;
        padding: 12px 14px;
        margin-bottom: 12px;
        font-size: 0.88rem;
        line-height: 1.5;
        backdrop-filter: blur(4px);
    }
    .workflow-step { 
        font-weight: 800; 
        color: #38bdf8; 
        margin-bottom: 4px; 
        font-size: 0.95rem;
        text-transform: uppercase;
        letter-spacing: 0.05em;
    }
    
    /* Layout des Tabs et TextAreas */
    .soft-panel {
        background: rgba(255,255,255,0.8);
        border: 1px solid rgba(148,163,184,0.18);
        border-radius: 20px;
        padding: 1.5rem;
        box-shadow: 0 12px 34px rgba(15,23,42,0.06);
    }
    .stTextArea textarea {
        border-radius: 18px !important;
        border: 1px solid var(--border) !important;
        padding: 1rem !important;
    }
    </style>
    """, unsafe_allow_html=True)

# 3. Fonction pour formater les balises
def formater_texte_anonymise(texte):
    texte_html = texte.replace('\n', '<br>')
    pattern = r"(\[[A-Z0-9_\-:\s]+\])"
    remplacement = r"<span class='badge-anonyme'>\1</span>"
    return re.sub(pattern, remplacement, texte_html)

# 4. Générateurs visuels (PDF / Word)
def anonymiser_docx_visuel(fichier_upload, liste_pii):
    fichier_upload.seek(0)
    doc = docx.Document(fichier_upload)
    pii_tries = sorted(liste_pii, key=lambda x: len(x["text"]), reverse=True)

    def remplacer_dans_paragraphe(para, pii_list):
        for pii in pii_list:
            texte_cible = pii["text"]
            texte_remplacement = str(pii.get("replacement", "[MASQUÉ]"))
            if texte_cible in para.text:
                pattern = re.compile(rf'(?<![\wÀ-ÿ\[]){re.escape(texte_cible)}(?![\wÀ-ÿ\]])')
                remplace_dans_runs = False
                for run in para.runs:
                    if texte_cible in run.text and pattern.search(run.text):
                        run.text = pattern.sub(texte_remplacement, run.text)
                        remplace_dans_runs = True
                if not remplace_dans_runs and pattern.search(para.text):
                    para.text = pattern.sub(texte_remplacement, para.text)
        if "*" in para.text or "|" in para.text:
            for run in para.runs:
                if "*" in run.text or "|" in run.text:
                    run.text = run.text.replace("*", "").replace("|", "")

    def traiter_bloc_texte(blocs_paragraphes, blocs_tableaux):
        for para in blocs_paragraphes: remplacer_dans_paragraphe(para, pii_tries)
        for table in blocs_tableaux:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs: remplacer_dans_paragraphe(para, pii_tries)

    traiter_bloc_texte(doc.paragraphs, doc.tables)
    for section in doc.sections:
        traiter_bloc_texte(section.header.paragraphs, section.header.tables)
        traiter_bloc_texte(section.footer.paragraphs, section.footer.tables)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def anonymiser_pdf_visuel(fichier_upload, liste_pii):
    fichier_upload.seek(0)
    pdf_bytes = fichier_upload.getvalue()
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pii_tries = sorted(liste_pii, key=lambda x: len(x["text"]), reverse=True)

    for page in doc:
        for pii in pii_tries:
            texte_original = pii["text"]
            texte_remplacement = pii["replacement"]
            instances = page.search_for(texte_original)
            for inst in instances:
                page.add_redact_annot(inst, text=texte_remplacement, fill=(0.95, 0.95, 0.95), text_color=(0, 0, 0), fontsize=9)
        page.apply_redactions()
    return doc.write()

# 5. Initialisation des modules
@st.cache_resource
def load_agent(): return AnonymizationAgent()

@st.cache_resource
def load_pseudonymizer(): return PatientPseudonymizer()

try:
    agent = load_agent()
    pseudo_engine = load_pseudonymizer()
except Exception as e:
    st.error(f"Erreur d'initialisation du système : {e}")
    st.stop()

# 6. Barre latérale avec le WORKFLOW
with st.sidebar:
    st.markdown("## HOSIX")
    st.success("🟢 Connecté — environnement sécurisé")
    
    st.markdown("### ⚙️ Workflow (Mode d'emploi)")
    st.markdown("""
    <div class="workflow-box">
        <div class="workflow-step">1. Importer</div>
        Saisissez un texte, importez un document unique, ou sélectionnez un lot de fichiers.
    </div>
    <div class="workflow-box">
        <div class="workflow-step">2. Analyse IA (NLP)</div>
        L'agent identifie automatiquement les informations sensibles dans le texte.
    </div>
    <div class="workflow-box">
        <div class="workflow-step">3. Sécurisation</div>
        Génération des pseudonymes et chiffrement AES-128 dans la base de données.
    </div>
    <div class="workflow-box">
        <div class="workflow-step">4. Exportation</div>
        Téléchargez le document expurgé de toute donnée directe, prêt à l'emploi.
    </div>
    """, unsafe_allow_html=True)

# 7. En-tête (Hero Section MIS À JOUR)
st.markdown(
    """
    <div class="hero">
        <div class="hero-title">Pipeline de dé-identification clinique</div>
        <p class="hero-subtitle">
            Anonymisez vos comptes rendus médicaux avec une présentation claire, 
            élégante et conforme aux réglementations de santé en vigueur. 
            <strong>Cliquez sur les cartes ci-dessous pour afficher les détails.</strong>
        </p>
    </div>
    """, unsafe_allow_html=True
)

# 8. Cartes interactives (Popovers pour les Concepts et PII)
col_c1, col_c2, col_c3 = st.columns(3)
with col_c1:
    with st.popover("⚖️ Cadre Légal : Loi 09-08 & CNDP", use_container_width=True):
        st.markdown("### Protection des données au Maroc")
        st.write("Ce système est conçu en stricte conformité avec la **Loi n° 09-08** relative à la protection des personnes physiques à l'égard du traitement des données à caractère personnel. Il garantit un traitement licite et la confidentialité absolue du patient.")
with col_c2:
    with st.popover("🛡️ Sécurité : Zero-Trust & AES", use_container_width=True):
        st.markdown("### Architecture sécurisée")
        st.write("L'application applique une politique **Zero-Trust**. Les données directes sont pseudonymisées via une fonction de hachage robuste (HMAC-SHA256 couplé à un Pepper). Le stockage en base est chiffré via AES-128 (Fernet).")
with col_c3:
    with st.popover("📄 Formats & Traitement par Lots", use_container_width=True):
        st.markdown("### Flexibilité d'utilisation")
        st.write("Prise en charge native des fichiers **TXT, PDF et DOCX**. Vous pouvez traiter un compte rendu unique ou importer un lot de plusieurs dizaines de documents simultanément, tout en préservant la mise en page lors de l'export.")

col_p1, col_p2 = st.columns(2)
with col_p1:
    with st.popover("🔴 Ce que nous masquons (PII Identifiantes)", use_container_width=True):
        st.markdown("### Informations supprimées ou chiffrées :")
        st.markdown("""
        * Identité complète (Nom, Prénom)
        * Identifiants administratifs (CIN, N° Dossier, IP/IPP)
        * Coordonnées (Téléphone, Adresse exacte, Email)
        * Date de naissance exacte et identifiants de couverture
        """)
        st.info("💡 **Pseudonymisation :** L'identité est remplacée par une clé unique (ex: `[PAT-A1B2C3]`), permettant un suivi longitudinal clinique sans révéler le vrai nom.")
with col_p2:
    with st.popover("🟢 Ce que nous conservons (Épidémiologie)", use_container_width=True):
        st.markdown("### Informations utiles conservées en clair :")
        st.markdown("""
        * Ville ou Région géographique
        * Mois & Année de consultation (jour exact masqué)
        * Organisme de Couverture (ex: AMO, CNSS)
        * L'intégralité des données cliniques, diagnostics et traitements
        """)
        st.success("✅ **Résultat :** Le document reste parfaitement lisible et exploitable pour la recherche, l'apprentissage IA ou les statistiques hospitalières.")

st.markdown("<br>", unsafe_allow_html=True)

# 9. Extraction de texte
def extraire_texte(fichier):
    try:
        if fichier.name.endswith(".txt"): return fichier.getvalue().decode("utf-8")
        elif fichier.name.endswith(".docx"):
            doc = docx.Document(fichier)
            texte_complet = [para.text for para in doc.paragraphs if para.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    ligne = [cell.text.replace('\n', ' ').strip() for cell in row.cells if cell.text.strip()]
                    if ligne: texte_complet.append(" | ".join(ligne))
            return "\n".join(texte_complet)
        elif fichier.name.endswith(".pdf"):
            texte_final = ""
            with pdfplumber.open(fichier) as pdf:
                for page in pdf.pages:
                    texte = page.extract_text(layout=True)
                    if texte: texte_final += texte + "\n"
            return texte_final
    except Exception as e:
        st.error(f"Erreur de lecture : {e}")
        return ""


# 10. INTERFACE À 3 ONGLETS
tab1, tab2, tab3 = st.tabs(["✍️ Saisie manuelle", "📄 Document unique", "📚 Traitement par lots"])

texte_cible = ""
fichier_cible = None
fichiers_a_traiter = []
lancer_traitement_unitaire = False
lancer_traitement_lots = False

with tab1:
    st.markdown("##### Coller le texte à anonymiser")
    texte_manuel = st.text_area(
        label="Texte source", height=200, label_visibility="collapsed",
        placeholder="Ex. Le 12/05/2026, consultation du patient M. Ahmed..."
    )
    if st.button("Anonymiser le texte", key="btn_manuel"):
        if texte_manuel.strip():
            texte_cible = texte_manuel
            lancer_traitement_unitaire = True
        else: st.warning("Veuillez saisir un texte.")
    st.markdown('</div>', unsafe_allow_html=True)

with tab2:
    st.markdown("##### Importer un document spécifique")
    fichier_unique = st.file_uploader("Formats acceptés : TXT, PDF, DOCX", type=["txt", "pdf", "docx"], key="up_unique")
    if st.button("Anonymiser le document", key="btn_unique"):
        if fichier_unique:
            texte_cible = extraire_texte(fichier_unique)
            fichier_cible = fichier_unique
            lancer_traitement_unitaire = True
        else: st.warning("Veuillez sélectionner un fichier.")
    st.markdown('</div>', unsafe_allow_html=True)

with tab3:
    st.markdown("##### Traitement par lots")
    st.info("Importez plusieurs documents. Les résultats seront générés sous forme de liste téléchargeable.")
    fichiers_lots = st.file_uploader("Sélectionnez vos fichiers", type=["txt", "pdf", "docx"], accept_multiple_files=True, key="up_lots")
    if st.button("Lancer le traitement", key="btn_lots"):
        if fichiers_lots:
            fichiers_a_traiter = fichiers_lots
            lancer_traitement_lots = True
        else: st.warning("Veuillez sélectionner au moins un fichier.")
    st.markdown('</div>', unsafe_allow_html=True)

st.divider()

# ==========================================
# 11. LOGIQUE DE TRAITEMENT
# ==========================================

# --- FLUX A : TRAITEMENT UNITAIRE ---
if lancer_traitement_unitaire and texte_cible.strip():
    with st.status("🔍 Analyse syntaxique et IA en cours...", expanded=True) as status:
        try:
            resultat_nlp = agent.run(texte_cible)
            status.update(label="🛡️ Chiffrement et génération du pseudonyme...")
            
            liste_pii = resultat_nlp["llm_json"].get("pii_detected", [])
            for p in liste_pii:
                if "*" in p["text"]: p["type"] = "IP"  
            
            ip_match = re.search(r"(?:IP|IPP)[\s:\-\n]*([*]?[\d\s]{5,15}[*]?)", texte_cible, re.IGNORECASE)
            valeur_ip_pure = ""
            if ip_match:
                valeur_ip_brute = ip_match.group(1).strip() 
                valeur_ip_pure = valeur_ip_brute.replace('*', '').strip() 
                if not any(p.get("type") in ["IP", "IPP", "OTHER_ID"] for p in liste_pii):
                    liste_pii.append({"text": valeur_ip_brute, "type": "IP", "replacement": ""})
                morceaux = valeur_ip_pure.split()
                if len(morceaux) > 1:
                    for morceau in morceaux:
                        if len(morceau) >= 3: liste_pii.append({"text": morceau, "type": "IP", "replacement": ""})

            cle_patient = pseudo_engine.construire_cle_patient(liste_pii)
            hash_patient, pseudonyme_id, db_status = pseudo_engine.generer_pseudonyme(cle_patient, liste_pii)
            
            status.update(label="✅ Anonymisation terminée avec succès !", state="complete", expanded=False)

            # Attribution des remplacements
            for pii in liste_pii:
                t = pii.get("type", "").upper()
                if t in ["PATIENT_NAME", "PERSON", "NOM_PATIENT"]:
                    if pseudonyme_id: pii["replacement"] = f"[{pseudonyme_id}]"
                elif t in ["DOSSIER_MEDICAL", "DOSSIER_MED"]: pii["replacement"] = "[DOSSIER_MEDICAL]"
                elif t in ["OTHER_ID", "IP", "IPP"]: pii["replacement"] = ""

            texte_final_brut = resultat_nlp["texte_anonymise_llm"]
            if pseudonyme_id:
                texte_final_brut = re.sub(r"\[PATIENT_NAME\]|\[PATIENT\]|\[PERSONNE\]|\[NOM_PATIENT\]", f"[{pseudonyme_id}]", texte_final_brut)
                texte_final_brut = re.sub(rf"[\(\[]{pseudonyme_id}[\)\]](?:\s*[\(\[]{pseudonyme_id}[\)\]])+", f"[{pseudonyme_id}]", texte_final_brut)
            
            texte_final_brut = re.sub(r"\[IP\]|\[IPP\]|\[OTHER_ID\]", "", texte_final_brut)
            if valeur_ip_pure: texte_final_brut = re.sub(re.escape(valeur_ip_pure), "", texte_final_brut)
            texte_final_brut = re.sub(r"\[DOSSIER_MED\]", "[DOSSIER_MEDICAL]", texte_final_brut)

            # Affichage visuel conditionnel
            if fichier_cible is None or fichier_cible.name.endswith(".txt"):
                col1, col2 = st.columns(2)
                with col1:
                    st.markdown("#### 📄 Texte Original")
                    st.markdown(f'<div class="text-card">{texte_cible.replace(chr(10), "<br>")}</div>', unsafe_allow_html=True)
                with col2:
                    st.markdown("#### 🛡️ Texte Pseudonymisé")
                    st.markdown(f'<div class="text-card">{formater_texte_anonymise(texte_final_brut)}</div>', unsafe_allow_html=True)
            else:
                st.info("📄 Téléchargez le fichier ci-dessous pour voir le résultat complet.")
            
            st.markdown("<br>", unsafe_allow_html=True)
            if fichier_cible and fichier_cible.name.endswith(".pdf"):
                st.download_button("📥 Télécharger PDF Sécurisé", data=anonymiser_pdf_visuel(fichier_cible, liste_pii), file_name="CR_anonymise.pdf", mime="application/pdf")
            elif fichier_cible and fichier_cible.name.endswith(".docx"):
                st.download_button("📥 Télécharger Word Sécurisé", data=anonymiser_docx_visuel(fichier_cible, liste_pii), file_name="CR_anonymise.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            else:
                st.download_button("📥 Télécharger TXT Sécurisé", data=texte_final_brut, file_name="CR_anonymise.txt", mime="text/plain")

        except Exception as e:
            st.error(f"⚠️ Erreur lors du traitement : {e}")

# --- FLUX B : TRAITEMENT PAR LOTS ---
elif lancer_traitement_lots and fichiers_a_traiter:
    st.markdown("### ⚙️ Progression du traitement")
    barre = st.progress(0)
    texte_statut = st.empty()
    zone_resultats = st.container()
    
    for index, fichier in enumerate(fichiers_a_traiter):
        texte_statut.text(f"Traitement en cours : {fichier.name} ({index + 1}/{len(fichiers_a_traiter)})")
        texte_cible_lot = extraire_texte(fichier)
        
        if texte_cible_lot.strip():
            try:
                resultat_nlp = agent.run(texte_cible_lot)
                liste_pii = resultat_nlp["llm_json"].get("pii_detected", [])
                
                # Gestion des IP et code-barres
                for p in liste_pii:
                    if "*" in p["text"]: p["type"] = "IP"  
                
                ip_match = re.search(r"(?:IP|IPP)[\s:\-\n]*([*]?[\d\s]{5,15}[*]?)", texte_cible_lot, re.IGNORECASE)
                valeur_ip_pure = ""
                if ip_match:
                    valeur_ip_brute = ip_match.group(1).strip()
                    valeur_ip_pure = valeur_ip_brute.replace('*', '').strip() 
                    liste_pii.append({"text": valeur_ip_brute, "type": "IP", "replacement": ""})
                    liste_pii.append({"text": valeur_ip_pure, "type": "IP", "replacement": ""})
                    morceaux = valeur_ip_pure.split()
                    if len(morceaux) > 1:
                        for morceau in morceaux:
                            if len(morceau) >= 3: liste_pii.append({"text": morceau, "type": "IP", "replacement": ""})
                
                # Base de données
                cle_patient = pseudo_engine.construire_cle_patient(liste_pii)
                hash_patient, pseudonyme_id, db_status = pseudo_engine.generer_pseudonyme(cle_patient, liste_pii)
                
                # Attribution des remplacements PII pour les lots
                for pii in liste_pii:
                    t = pii.get("type", "").upper()
                    if t in ["PATIENT_NAME", "PERSON", "NOM_PATIENT"]:
                        if pseudonyme_id: pii["replacement"] = f"[{pseudonyme_id}]"
                    elif t in ["DOSSIER_MEDICAL", "DOSSIER_MED"]: pii["replacement"] = "[DOSSIER_MEDICAL]"
                    elif t in ["OTHER_ID", "IP", "IPP"]: pii["replacement"] = ""

                # Remplacement brut pour le TXT
                texte_final_brut = resultat_nlp["texte_anonymise_llm"]
                if pseudonyme_id:
                    texte_final_brut = re.sub(r"\[PATIENT_NAME\]|\[PATIENT\]|\[PERSONNE\]|\[NOM_PATIENT\]", f"[{pseudonyme_id}]", texte_final_brut)
                    texte_final_brut = re.sub(rf"[\(\[]{pseudonyme_id}[\)\]](?:\s*[\(\[]{pseudonyme_id}[\)\]])+", f"[{pseudonyme_id}]", texte_final_brut)
                texte_final_brut = re.sub(r"\[IP\]|\[IPP\]|\[OTHER_ID\]", "", texte_final_brut)
                if valeur_ip_pure: texte_final_brut = re.sub(re.escape(valeur_ip_pure), "", texte_final_brut)
                texte_final_brut = re.sub(r"\[DOSSIER_MED\]", "[DOSSIER_MEDICAL]", texte_final_brut)

                with zone_resultats:
                    with st.expander(f"✅ Terminé : {fichier.name} | Pseudo : {pseudonyme_id or 'Non assigné'}"):
                        if fichier.name.endswith(".pdf"):
                            st.download_button("📥 Télécharger PDF", data=anonymiser_pdf_visuel(fichier, liste_pii), file_name=f"Anon_{fichier.name}", mime="application/pdf", key=f"dl_{index}")
                        elif fichier.name.endswith(".docx"):
                            st.download_button("📥 Télécharger Word", data=anonymiser_docx_visuel(fichier, liste_pii), file_name=f"Anon_{fichier.name}", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"dl_{index}")
                        else:
                            st.download_button("📥 Télécharger TXT", data=texte_final_brut, file_name=f"Anon_{fichier.name}", mime="text/plain", key=f"dl_{index}")
            except Exception as e:
                with zone_resultats:
                    st.error(f"❌ Échec pour {fichier.name}")
        
        barre.progress((index + 1) / len(fichiers_a_traiter))
        
    texte_statut.success("🎉 Traitement par lots terminé avec succès !")